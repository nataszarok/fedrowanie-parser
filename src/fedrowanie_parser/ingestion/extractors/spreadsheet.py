from __future__ import annotations
from pathlib import Path
import subprocess
import tempfile
from openpyxl import load_workbook
from ..models import CellTable
from ..runtime import find_libreoffice


def _xlsx_tables(path: Path, original: Path | None = None) -> list[CellTable]:
    wb=load_workbook(path, data_only=True, read_only=True)
    source=original or path
    out=[]
    for ws in wb.worksheets:
        rows=[]
        for row in ws.iter_rows(values_only=True):
            vals=list(row)
            while vals and vals[-1] is None:
                vals.pop()
            if vals and any(v not in (None, "") for v in vals):
                rows.append(vals)
        if rows:
            out.append(CellTable(source, ws.title, None, rows, "openpyxl"))
    return out


def extract_spreadsheet(path: Path) -> list[CellTable]:
    ext=path.suffix.lower()
    if ext==".xlsx":
        return _xlsx_tables(path)
    if ext==".xls":
        office=find_libreoffice()
        if not office:
            raise RuntimeError("legacy .xls requires LibreOffice/soffice; on macOS install with: brew install --cask libreoffice")
        with tempfile.TemporaryDirectory(prefix="fedrowanie-xls-") as td:
            subprocess.run([office,"--headless","--convert-to","xlsx","--outdir",td,str(path)], check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=60)
            converted=Path(td)/(path.stem+".xlsx")
            if not converted.exists():
                candidates=list(Path(td).glob("*.xlsx"))
                if not candidates:
                    raise RuntimeError("LibreOffice did not produce xlsx")
                converted=candidates[0]
            return _xlsx_tables(converted, original=path)
    raise ValueError(path)
